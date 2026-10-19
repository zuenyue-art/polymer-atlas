#!/usr/bin/env python3
"""把 props-待填.xlsx 里填好的数值合回 data.json。

三条刻意的规矩：

1. **留空 = 不改动**。已经填过的值不会因为这次留空被清掉。填表是分批做的，
   一次只填 30 条就把另外 100 条清掉是灾难；要改某个值就直接写新值。
2. **可疑值先报出来，不写进去**。数值超出该字段的合理区间（比如 Tg 填了
   900）多半是串行了或单位错了，静默入库比留空更糟——留空只是少一条数据，
   错的数会被别人当依据用。
3. **改完整体重建**。props 会进预渲染页和 Excel，所以合完之后要跑
   npm run build，脚本会在结尾提醒。

用法：
    python scripts/apply-props.py                  # 读 props-待填.xlsx
    python scripts/apply-props.py 我的表.xlsx
    python scripts/apply-props.py --dry-run        # 只报会改什么，不写盘
"""

import json
import os
import sys

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("缺少 openpyxl。请先安装：pip install -r scripts/requirements.txt")

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(ROOT, 'dist', 'data.json')
DEFAULT = os.path.join(ROOT, 'props-待填.xlsx')

# 合理区间。宽到不会误伤真实数据，窄到能拦住串行/写错单位。
BOUNDS = {
    'Tg': (-120, 400),
    'Tm': (0, 450),
    'HDT': (0, 400),
    '密度': (0.3, 3.5),
    '加工温度': (0, 500),
}


def to_num(v):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace('℃', '').replace('度', '')
    try:
        return float(s)
    except ValueError:
        return 'BAD'


def main():
    args = [a for a in sys.argv[1:] if not a.startswith('--')]
    dry = '--dry-run' in sys.argv
    src = args[0] if args else DEFAULT
    if not os.path.exists(src):
        sys.exit(f'找不到 {src}。先跑 python scripts/build-props-sheet.py 生成表格')

    data = json.load(open(DATA, encoding='utf-8'))
    spec = data.get('propsSpec') or []
    if not spec:
        sys.exit('data.json 里没有 propsSpec')
    by_id = {m['id']: m for m in data['materials']}

    wb = load_workbook(src, data_only=True)
    if '材料数值' not in wb.sheetnames:
        sys.exit(f'{src} 里没有「材料数值」工作表，这份表不对')
    ws = wb['材料数值']

    header = [c.value for c in ws[1]]
    # 列位置：先找 id 列，其余按 propsSpec 顺序推
    try:
        id_col = header.index('id')
    except ValueError:
        sys.exit('表头里没有 id 列')
    cols, i = {}, 4  # 前四列是 id / 分类 / 名称 / 缩写
    for s in spec:
        if s['type'] == 'range':
            cols[s['key']] = (i, i + 1)
            i += 2
        else:
            cols[s['key']] = (i, None)
            i += 1

    written, skipped, touched = 0, [], set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        mid = row[id_col]
        if not mid or mid not in by_id:
            continue
        m = by_id[mid]
        props = dict(m.get('props') or {})
        for s in spec:
            lo_i, hi_i = cols[s['key']]
            if s['type'] == 'range':
                lo, hi = to_num(row[lo_i]), to_num(row[hi_i])
                if lo is None and hi is None:
                    continue
                if lo == 'BAD' or hi == 'BAD':
                    skipped.append(f'{mid} {s["label"]}：不是数字')
                    continue
                lo = lo if lo is not None else hi
                hi = hi if hi is not None else lo
                b = BOUNDS.get(s['key'])
                if b and not (b[0] <= lo <= b[1] and b[0] <= hi <= b[1]):
                    skipped.append(f'{mid} {s["label"]}：{lo}~{hi} 超出合理区间 {b}')
                    continue
                if lo > hi:
                    lo, hi = hi, lo
                props[s['key']] = [lo, hi]
            else:
                v = to_num(row[lo_i])
                if v is None:
                    continue
                if v == 'BAD':
                    skipped.append(f'{mid} {s["label"]}：不是数字')
                    continue
                b = BOUNDS.get(s['key'])
                if b and not (b[0] <= v <= b[1]):
                    skipped.append(f'{mid} {s["label"]}：{v} 超出合理区间 {b}')
                    continue
                props[s['key']] = v
            written += 1
        if props != (m.get('props') or {}):
            m['props'] = props
            touched.add(mid)

    total = sum(1 for m in data['materials'] if m.get('props'))
    print(f'读取 {src}')
    print(f'  写入 {written} 个数值，涉及 {len(touched)} 条材料')
    print(f'  合完之后 {total} / {len(data["materials"])} 条材料有数值')

    if skipped:
        print(f'\n以下 {len(skipped)} 处没写进去（多半是串行或单位错，请回表里核对）：')
        for s in skipped[:20]:
            print('  ·', s)
        if len(skipped) > 20:
            print(f'  · 其余 {len(skipped) - 20} 处略')

    if dry:
        print('\n--dry-run：没有写盘。')
        return 0
    if not touched:
        print('\n没有变化，data.json 未改动。')
        return 0

    with open(DATA, 'w', encoding='utf-8', newline='\n') as f:
        f.write(json.dumps(data, ensure_ascii=False))
    print(f'\n已写入 {DATA}')
    print('接下来跑 npm run build —— props 会进预渲染页和可下载的 Excel。')
    return 0


if __name__ == '__main__':
    sys.exit(main())
