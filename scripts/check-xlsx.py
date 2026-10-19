"""校验 dist/materials.xlsx 与 data.json / terms.json 是否一致。

为什么用「读回来比对」而不是像 news 那样用 git diff 比对字节：xlsx 是 zip
包，里面每个条目都带写入时间戳，同样的内容两次生成字节必然不同。拿字节
做断言只会得到一堆假失败。这里比的是单元格里的值，那才是读者真正会看到
的东西。

用法：
    python scripts/check-xlsx.py
"""

import json
import os
import sys

try:
    from openpyxl import load_workbook
except ModuleNotFoundError:
    sys.exit("缺少 openpyxl。请先安装：pip install -r scripts/requirements.txt")

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DIST = os.path.join(ROOT, "dist")
XLSX = os.path.join(DIST, "materials.xlsx")

TERM_SHEET = "基础术语与使用说明"
TERM_COLUMNS = ["术语/阅读规则", "解释", "用途或易错点"]


def read_json(name):
    with open(os.path.join(DIST, name), encoding="utf-8") as f:
        return json.load(f)


def cell(value):
    """表格里的数字、None、字符串统一成字符串再比，避免类型差异造成假失败。"""
    return "" if value is None else str(value).strip()


def compare_sheet(problems, wb, title, columns, rows, label):
    if title not in wb.sheetnames:
        problems.append(f"{label}：缺少工作表「{title}」")
        return
    ws = wb[title]
    header = [cell(c) for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True), [])]
    if header != columns:
        problems.append(f"{label}：表头不一致\n    文件：{header}\n    数据：{columns}")
        return
    actual = [
        [cell(c) for c in r]
        for r in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True)
        if any(c is not None for c in r)
    ]
    expected = [[cell(v) for v in row] for row in rows]
    if len(actual) != len(expected):
        problems.append(f"{label}：行数不一致，文件 {len(actual)} 行 / 数据 {len(expected)} 行")
        return
    for i, (a, e) in enumerate(zip(actual, expected), start=2):
        if a != e:
            diff = [f"第 {j + 1} 列 文件={x!r} 数据={y!r}" for j, (x, y) in enumerate(zip(a, e)) if x != y]
            problems.append(f"{label}：第 {i} 行不一致（{'; '.join(diff[:3])}）")
            break


def main():
    if not os.path.exists(XLSX):
        print(f"缺少 {os.path.relpath(XLSX, ROOT)}，请先运行 npm run build:xlsx")
        return 1

    data = read_json("data.json")
    knowledge = read_json("terms.json")
    wb = load_workbook(XLSX, read_only=True, data_only=True)
    problems = []

    materials = data["materials"]
    columns = list(materials[0]["fields"].keys())
    categories = []
    for m in materials:
        if m["category"] not in categories:
            categories.append(m["category"])

    for category in categories:
        rows = [[m["fields"][c] for c in columns] for m in materials if m["category"] == category]
        compare_sheet(problems, wb, category, columns, rows, f"材料·{category}")

    terms = knowledge["terms"]
    compare_sheet(
        problems,
        wb,
        TERM_SHEET,
        TERM_COLUMNS,
        [[t[c] for c in TERM_COLUMNS] for t in terms],
        "术语",
    )

    extra = [s for s in wb.sheetnames if s not in categories and s != TERM_SHEET]
    if extra:
        problems.append(f"多出未预期的工作表：{extra}")

    if problems:
        print("materials.xlsx 与数据不一致：")
        for p in problems:
            print(f"  - {p}")
        print("\n修复：npm run build:xlsx")
        return 1

    print(
        f"materials.xlsx 与数据一致（{len(materials)} 条材料 / {len(terms)} 条术语）"
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
