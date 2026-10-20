"""从 data.json / terms.json 生成 dist/materials.xlsx。

为什么：首页有「下载完整 Excel」的入口，而这份文件此前是手工维护的第二份
数据副本。它不受任何检查约束，改了 JSON 忘了改表，读者下载到的就是过期
内容，而且没人会发现。现在它由数据生成——单一事实来源，不可能不一致。

术语取 terms.json（68 条正式词条）而不是 data.json 内嵌的那份：后者只是
terms.json 加载失败时的回退数据，条数少且会过时。站点上展示的是前者，
下载到的也应该是前者。

用法：
    python scripts/build-xlsx.py
"""

import json
import os
import sys

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
except ModuleNotFoundError:
    sys.exit("缺少 openpyxl。请先安装：pip install -r scripts/requirements.txt")

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DIST = os.path.join(ROOT, "dist")
OUT = os.path.join(DIST, "materials.xlsx")

# 表头样式沿用此前手工维护时的外观：深色底白字加粗，首行加高，冻结前两列。
HEAD_FILL = PatternFill("solid", fgColor="FF333333")
HEAD_FONT = Font(bold=True, color="FFFFFFFF")
HEAD_HEIGHT = 36
FREEZE = "C2"

# 列宽同样沿用旧文件的取值，避免读者下载后排版突变。
# 末尾 5 列是数值（Tg / Tm / HDT / 密度 / 加工温度），与 propsSpec 顺序一致。
MATERIAL_WIDTHS = [24, 20, 42, 27, 43, 40, 42, 40, 45, 45, 43, 52, 48, 72, 11, 11, 11, 12, 16]
TERM_WIDTHS = [29, 76, 74]

TERM_SHEET = "基础术语与使用说明"
# 术语表只写这三列：与旧文件的列结构保持一致，"分类"用于站点上的筛选，
# 放进表格反而会让读者以为它属于词条本身。
TERM_COLUMNS = ["术语/阅读规则", "解释", "用途或易错点"]


def read_json(name):
    with open(os.path.join(DIST, name), encoding="utf-8") as f:
        return json.load(f)


def write_sheet(wb, title, columns, rows, widths, first=False):
    if first:
        ws = wb.active
        ws.title = title
    else:
        ws = wb.create_sheet(title)
    ws.append(columns)
    for row in rows:
        ws.append(row)
    for cell in ws[1]:
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
    ws.row_dimensions[1].height = HEAD_HEIGHT
    ws.freeze_panes = FREEZE
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    return ws


def prop_columns(data):
    """数值列。字段定义只有 propsSpec 一处，表里不再写第二份。"""
    return [f'{s["label"]}({s["unit"]})' for s in (data.get("propsSpec") or [])]


def num(value):
    """58.0 写成 58：表格里不该出现「160.0 ~ 240.0」这种加工温度。"""
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def prop_cells(m, spec):
    """没填的数值留空而不是填 0——留空表示「没数据」，0 会被当成测过。"""
    p = m.get("props") or {}
    out = []
    for s in spec:
        v = p.get(s["key"])
        if isinstance(v, list):
            out.append(f"{num(v[0])} ~ {num(v[1])}")
        else:
            out.append(num(v))
    return out


def main():
    data = read_json("data.json")
    knowledge = read_json("terms.json")

    materials = data["materials"]
    spec = data.get("propsSpec") or []
    columns = list(materials[0]["fields"].keys()) + prop_columns(data)

    # 分类按首次出现顺序，与站点侧栏、sitemap 的顺序一致。
    order = []
    for m in materials:
        if m["category"] not in order:
            order.append(m["category"])

    wb = Workbook()
    for i, category in enumerate(order):
        rows = [
            [m["fields"][c] for c in materials[0]["fields"].keys()] + prop_cells(m, spec)
            for m in materials
            if m["category"] == category
        ]
        write_sheet(wb, category, columns, rows, MATERIAL_WIDTHS, first=i == 0)

    terms = knowledge["terms"]
    write_sheet(
        wb,
        TERM_SHEET,
        TERM_COLUMNS,
        [[t[c] for c in TERM_COLUMNS] for t in terms],
        TERM_WIDTHS,
    )

    wb.save(OUT)
    filled = sum(1 for m in materials for s in spec if (m.get("props") or {}).get(s["key"]) is not None)
    print(f"已生成 {os.path.relpath(OUT, ROOT)}")
    print(f"  材料 {len(materials)} 条 / {len(order)} 个分类工作表")
    print(f"  数值 {filled} 个（{len(spec)} 列，没填的留空）")
    print(f"  术语 {len(terms)} 条")
    return 0


if __name__ == "__main__":
    sys.exit(main())
